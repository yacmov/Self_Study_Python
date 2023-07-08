from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, LineChart, Reference
# # B2:C11 까지의 데이터를 차트로 생성
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart() # 차트 종류 설정 (Bar, line, pie,...)
# bar_chart.add_data(bar_value) # 차트 데이터 추가
# ws.add_chart(bar_chart, "E1") # E1에 차트를 넣어줌

# 제목을 포함한 데이터
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True) # 계열부분에 데이터 분류 이름이 들어간다
line_chart.title = "성적표" # 제목
line_chart.style = 15 # 미리 정의된 스타일을 적용
line_chart.y_axis.title = "점수" # 세로줄 이름 넣기
line_chart.x_axis.title = "번호"
ws.add_chart(line_chart, "E1")

wb.save("sample_chart.xlsx")